import streamlit as st
import pandas as pd
import io
import re

# ============================================================
# PAGE
# ============================================================

st.set_page_config(
    page_title="OBE Quiz Checker",
    page_icon="🎓",
    layout="wide"
)

# ============================================================
# OPTIONAL LIBRARIES
# ============================================================

try:
    import fitz
except Exception:
    fitz = None

try:
    import pypdf
except Exception:
    pypdf = None

try:
    import pdfplumber
except Exception:
    pdfplumber = None

try:
    from docx import Document
except Exception:
    Document = None

try:
    from pptx import Presentation
except Exception:
    Presentation = None

try:
    from PIL import Image
except Exception:
    Image = None

try:
    import pytesseract
except Exception:
    pytesseract = None

# ============================================================
# CONSTANTS
# ============================================================

ATTAINMENT = 75

BLOOM_LEVELS = [
    "Remember",
    "Understand",
    "Apply",
    "Analyze",
    "Evaluate",
    "Create"
]

BLOOM_VERBS = {
    "Remember": [
        "define", "identify", "list", "name", "state",
        "recall", "recognize", "label", "mention", "select"
    ],
    "Understand": [
        "describe", "explain", "summarize", "interpret",
        "classify", "discuss", "illustrate", "compare"
    ],
    "Apply": [
        "calculate", "solve", "apply", "demonstrate",
        "use", "implement", "compute", "perform"
    ],
    "Analyze": [
        "analyze", "analyse", "differentiate", "examine",
        "investigate", "contrast", "distinguish", "categorize"
    ],
    "Evaluate": [
        "evaluate", "justify", "assess", "critique",
        "judge", "defend", "recommend", "appraise"
    ],
    "Create": [
        "create", "design", "develop", "formulate",
        "produce", "construct", "propose", "generate", "plan"
    ]
}

STOPWORDS = {
    "the", "a", "an", "and", "or", "of", "to", "in", "on",
    "for", "from", "with", "by", "is", "are", "was", "were",
    "be", "been", "being", "this", "that", "these", "those",
    "it", "its", "their", "them", "they", "you", "your",
    "using", "use", "used", "student", "students", "course",
    "learners", "learner", "will", "can", "should", "may"
}

# ============================================================
# TEXT UTILITIES
# ============================================================

def clean_text(text):
    if text is None:
        return ""

    text = str(text)

    replacements = {
        "\x00": " ",
        "\uf0b7": "•",
        "\u00a0": " ",
        "\u2013": "-",
        "\u2014": "-",
        "\u2018": "'",
        "\u2019": "'",
        "\u201c": '"',
        "\u201d": '"'
    }

    for a, b in replacements.items():
        text = text.replace(a, b)

    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n[ \t]+", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)

    return text.strip()


def normalize(text):
    text = clean_text(text).lower()
    text = re.sub(r"[^a-z0-9\s]", " ", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def words(text):
    result = set()

    for word in normalize(text).split():

        if len(word) <= 2:
            continue

        if word in STOPWORDS:
            continue

        if word.endswith("ies") and len(word) > 4:
            word = word[:-3] + "y"
        elif word.endswith("ing") and len(word) > 5:
            word = word[:-3]
        elif word.endswith("ed") and len(word) > 4:
            word = word[:-2]
        elif word.endswith("s") and len(word) > 4:
            word = word[:-1]

        result.add(word)

    return result


def extract_numbers(text):
    return re.findall(
        r"\b\d+(?:\.\d+)?\b",
        clean_text(text)
    )


# ============================================================
# OUTCOMES
# ============================================================

def parse_outcomes(text):
    if not text:
        return []

    lines = re.split(r"[\n;]+", text)
    result = []

    for line in lines:

        line = clean_text(line)

        if not line:
            continue

        line = re.sub(
            r"^(?:CLO|PLO|CO|PO)\s*\d*\s*[:.)-]?\s*",
            "",
            line,
            flags=re.I
        )

        line = re.sub(
            r"^\d+\s*[\).:-]\s*",
            "",
            line
        )

        if len(line) >= 8:
            result.append(line)

    final = []
    seen = set()

    for item in result:
        key = normalize(item)

        if key not in seen:
            seen.add(key)
            final.append(item)

    return final


# ============================================================
# BLOOM
# ============================================================

def detect_bloom(text):

    t = normalize(text)

    for level in BLOOM_LEVELS:

        for verb in BLOOM_VERBS[level]:

            if re.search(
                r"\b" + re.escape(verb) + r"\b",
                t
            ):
                return level

    return "Understand"


def bloom_score(question, target):

    actual = detect_bloom(question)

    try:
        a = BLOOM_LEVELS.index(actual)
        b = BLOOM_LEVELS.index(target)
    except Exception:
        return 80

    distance = abs(a - b)

    if distance == 0:
        return 100

    if distance == 1:
        return 92

    if distance == 2:
        return 84

    if distance == 3:
        return 78

    return 72


# ============================================================
# PDF READING
# ============================================================

def read_pdf(uploaded_file):

    data = uploaded_file.getvalue()

    if not data:
        raise ValueError("The uploaded PDF is empty.")

    text_parts = []

    # --------------------------------------------------------
    # METHOD 1: PyMuPDF
    # --------------------------------------------------------

    if fitz is not None:

        try:

            doc = fitz.open(
                stream=data,
                filetype="pdf"
            )

            for page_no, page in enumerate(
                doc,
                start=1
            ):

                text = page.get_text(
                    "text",
                    sort=True
                )

                text = clean_text(text)

                if text:
                    text_parts.append(
                        f"--- PAGE {page_no} ---\n{text}"
                    )

            doc.close()

        except Exception:
            pass

    # --------------------------------------------------------
    # METHOD 2: pypdf
    # --------------------------------------------------------

    if not text_parts and pypdf is not None:

        try:

            reader = pypdf.PdfReader(
                io.BytesIO(data)
            )

            for page_no, page in enumerate(
                reader.pages,
                start=1
            ):

                text = page.extract_text()

                text = clean_text(text)

                if text:
                    text_parts.append(
                        f"--- PAGE {page_no} ---\n{text}"
                    )

        except Exception:
            pass

    # --------------------------------------------------------
    # METHOD 3: pdfplumber
    # --------------------------------------------------------

    if not text_parts and pdfplumber is not None:

        try:

            with pdfplumber.open(
                io.BytesIO(data)
            ) as pdf:

                for page_no, page in enumerate(
                    pdf.pages,
                    start=1
                ):

                    text = page.extract_text()

                    text = clean_text(text)

                    if text:
                        text_parts.append(
                            f"--- PAGE {page_no} ---\n{text}"
                        )

        except Exception:
            pass

    extracted = clean_text(
        "\n\n".join(text_parts)
    )

    # --------------------------------------------------------
    # METHOD 4: OCR
    # --------------------------------------------------------

    if len(extracted) < 50:

        if (
            fitz is not None
            and pytesseract is not None
            and Image is not None
        ):

            try:

                doc = fitz.open(
                    stream=data,
                    filetype="pdf"
                )

                ocr_parts = []

                for page_no, page in enumerate(
                    doc,
                    start=1
                ):

                    pix = page.get_pixmap(
                        matrix=fitz.Matrix(
                            2.5,
                            2.5
                        ),
                        alpha=False
                    )

                    image = Image.open(
                        io.BytesIO(
                            pix.tobytes("png")
                        )
                    )

                    text = pytesseract.image_to_string(
                        image
                    )

                    text = clean_text(text)

                    if text:
                        ocr_parts.append(
                            f"--- PAGE {page_no} ---\n{text}"
                        )

                doc.close()

                ocr_text = clean_text(
                    "\n\n".join(ocr_parts)
                )

                if len(ocr_text) > len(extracted):
                    extracted = ocr_text

            except Exception:
                pass

    if not extracted:
        raise ValueError(
            "The PDF was opened, but no readable text could be extracted."
        )

    return extracted


# ============================================================
# DOCX
# ============================================================

def read_docx(uploaded_file):

    if Document is None:
        raise ValueError(
            "python-docx is not installed."
        )

    document = Document(
        io.BytesIO(
            uploaded_file.getvalue()
        )
    )

    parts = []

    for paragraph in document.paragraphs:

        text = clean_text(
            paragraph.text
        )

        if text:
            parts.append(text)

    for table in document.tables:

        for row in table.rows:

            values = []

            for cell in row.cells:

                value = clean_text(
                    cell.text
                )

                if value:
                    values.append(value)

            if values:
                parts.append(
                    " | ".join(values)
                )

    return clean_text(
        "\n".join(parts)
    )


# ============================================================
# PPTX
# ============================================================

def read_pptx(uploaded_file):

    if Presentation is None:
        raise ValueError(
            "python-pptx is not installed."
        )

    presentation = Presentation(
        io.BytesIO(
            uploaded_file.getvalue()
        )
    )

    parts = []

    for slide_no, slide in enumerate(
        presentation.slides,
        start=1
    ):

        parts.append(
            f"--- SLIDE {slide_no} ---"
        )

        for shape in slide.shapes:

            if hasattr(shape, "text"):

                text = clean_text(
                    shape.text
                )

                if text:
                    parts.append(text)

    return clean_text(
        "\n".join(parts)
    )


# ============================================================
# EXCEL
# ============================================================

def read_excel(uploaded_file):

    data = uploaded_file.getvalue()

    parts = []

    try:

        sheets = pd.read_excel(
            io.BytesIO(data),
            sheet_name=None,
            header=None
        )

        for sheet_name, df in sheets.items():

            parts.append(
                f"--- SHEET {sheet_name} ---"
            )

            df = df.fillna("")

            for row in df.astype(str).values:

                values = []

                for value in row:

                    value = clean_text(
                        value
                    )

                    if value:
                        values.append(value)

                if values:
                    parts.append(
                        " | ".join(values)
                    )

    except Exception as exc:

        raise ValueError(
            f"Excel file could not be read: {exc}"
        )

    return clean_text(
        "\n".join(parts)
    )


# ============================================================
# CSV
# ============================================================

def read_csv(uploaded_file):

    data = uploaded_file.getvalue()

    try:

        df = pd.read_csv(
            io.BytesIO(data),
            header=None
        )

    except Exception:

        df = pd.read_csv(
            io.BytesIO(data),
            header=None,
            encoding="latin-1"
        )

    parts = []

    for row in df.fillna("").astype(str).values:

        values = []

        for value in row:

            value = clean_text(value)

            if value:
                values.append(value)

        if values:
            parts.append(
                " | ".join(values)
            )

    return clean_text(
        "\n".join(parts)
    )


# ============================================================
# TEXT FILE
# ============================================================

def read_text(uploaded_file):

    data = uploaded_file.getvalue()

    try:
        text = data.decode(
            "utf-8"
        )
    except Exception:
        text = data.decode(
            "latin-1"
        )

    return clean_text(text)


# ============================================================
# IMAGE
# ============================================================

def read_image(uploaded_file):

    if Image is None:
        raise ValueError(
            "Pillow is not installed."
        )

    if pytesseract is None:
        raise ValueError(
            "pytesseract is not installed."
        )

    image = Image.open(
        io.BytesIO(
            uploaded_file.getvalue()
        )
    )

    text = pytesseract.image_to_string(
        image
    )

    return clean_text(text)


# ============================================================
# UNIVERSAL FILE READER
# ============================================================

def read_file(uploaded_file):

    name = uploaded_file.name.lower()

    if name.endswith(".pdf"):
        return read_pdf(uploaded_file)

    if name.endswith(".docx"):
        return read_docx(uploaded_file)

    if name.endswith(".pptx"):
        return read_pptx(uploaded_file)

    if name.endswith(
        (".xlsx", ".xls")
    ):
        return read_excel(uploaded_file)

    if name.endswith(".csv"):
        return read_csv(uploaded_file)

    if name.endswith(
        (".txt", ".md", ".text")
    ):
        return read_text(uploaded_file)

    if name.endswith(
        (
            ".png",
            ".jpg",
            ".jpeg",
            ".webp",
            ".bmp",
            ".tif",
            ".tiff"
        )
    ):
        return read_image(uploaded_file)

    raise ValueError(
        "Unsupported file type."
    )


# ============================================================
# QUESTION EXTRACTION
# ============================================================

def extract_questions(text):

    text = clean_text(text)

    if not text:
        return []

    questions = []

    # --------------------------------------------------------
    # METHOD 1: Numbered questions
    # --------------------------------------------------------

    pattern = re.compile(
        r"(?=(?:^|\n|\s)"
        r"(?:question\s*)?"
        r"(?:q\s*)?"
        r"\d{1,3}"
        r"\s*[\.\):\-]\s+)",
        re.I
    )

    chunks = re.split(
        pattern,
        text
    )

    for chunk in chunks:

        chunk = clean_text(chunk)

        chunk = re.sub(
            r"^(?:question\s*)?"
            r"(?:q\s*)?"
            r"\d{1,3}"
            r"\s*[\.\):\-]\s*",
            "",
            chunk,
            flags=re.I
        )

        if len(chunk) >= 10:
            questions.append(chunk)

    # --------------------------------------------------------
    # METHOD 2: Question marks
    # --------------------------------------------------------

    if not questions:

        parts = re.split(
            r"(?<=[?])\s+",
            text
        )

        for part in parts:

            part = clean_text(part)

            if len(part) >= 10 and "?" in part:
                questions.append(part)

    # --------------------------------------------------------
    # METHOD 3: Question verbs at line starts
    # --------------------------------------------------------

    if not questions:

        lines = text.splitlines()

        buffer = ""

        verbs = []

        for values in BLOOM_VERBS.values():
            verbs.extend(values)

        verbs.extend([
            "what",
            "why",
            "how",
            "when",
            "where",
            "which",
            "who",
            "write",
            "give"
        ])

        for line in lines:

            line = clean_text(line)

            if not line:
                continue

            low = normalize(line)

            starts_question = any(
                low.startswith(
                    verb + " "
                )
                or low == verb
                for verb in verbs
            )

            if starts_question:

                if buffer:
                    questions.append(
                        clean_text(buffer)
                    )
                    buffer = ""

                questions.append(line)

            else:

                if buffer:
                    buffer += " " + line
                else:
                    buffer = line

        if buffer and len(buffer) >= 10:
            questions.append(
                clean_text(buffer)
            )

    # --------------------------------------------------------
    # METHOD 4: Paragraph fallback
    # --------------------------------------------------------

    if not questions:

        paragraphs = re.split(
            r"\n\s*\n",
            text
        )

        for paragraph in paragraphs:

            paragraph = clean_text(
                paragraph
            )

            if len(paragraph) >= 15:
                questions.append(
                    paragraph
                )

    # --------------------------------------------------------
    # Clean
    # --------------------------------------------------------

    final = []
    seen = set()

    for question in questions:

        question = clean_text(
            question
        )

        question = re.sub(
            r"---\s*(PAGE|SLIDE|SHEET).*?---",
            "",
            question,
            flags=re.I
        )

        question = clean_text(
            question
        )

        if len(question) < 10:
            continue

        # Ignore headings
        if normalize(question) in {
            "quiz",
            "assessment",
            "questions",
            "quiz questions",
            "instructions",
            "assignment",
            "assignment questions"
        }:
            continue

        key = normalize(question)

        if key not in seen:
            seen.add(key)
            final.append(question)

    return final


# ============================================================
# QUESTION TYPE
# ============================================================

def get_question_type(question):

    q = clean_text(question)

    options = re.findall(
        r"(?:^|\n|\s)[A-Da-d][\)\.\:]\s+",
        q
    )

    if len(options) >= 2:
        return "MCQ"

    if re.search(
        r"\btrue\s*/\s*false\b",
        q,
        re.I
    ):
        return "True / False"

    if "_" in q:
        return "Fill in the Blank"

    if re.search(
        r"\b(calculate|compute|solve|derive)\b",
        q,
        re.I
    ) and re.search(
        r"\d",
        q
    ):
        return "Numerical / Problem Solving"

    if re.search(
        r"\b(case study|scenario)\b",
        q,
        re.I
    ):
        return "Case Study / Application"

    if len(q.split()) > 45:
        return "Essay / Long Answer"

    return "Short Answer"


# ============================================================
# MATCHING
# ============================================================

def match_score(question, outcome):

    q = words(question)
    o = words(outcome)

    if not q or not o:
        return 55

    overlap = q.intersection(o)

    ratio = len(overlap) / max(
        1,
        len(o)
    )

    if ratio >= 0.70:
        return 95

    if ratio >= 0.50:
        return 88

    if ratio >= 0.35:
        return 78

    if ratio >= 0.20:
        return 68

    if ratio > 0:
        return 60

    return 55


def best_match(question, outcomes):

    if not outcomes:
        return "", 80

    scored = []

    for outcome in outcomes:

        scored.append(
            (
                match_score(
                    question,
                    outcome
                ),
                outcome
            )
        )

    scored.sort(
        reverse=True
    )

    return scored[0][1], scored[0][0]


# ============================================================
# OTHER METRICS
# ============================================================

def clarity(question):

    score = 98

    if len(question.split()) > 60:
        score -= 6

    elif len(question.split()) > 45:
        score -= 3

    if "etc." in question.lower():
        score -= 3

    if re.search(
        r"\bthing|things|stuff\b",
        question,
        re.I
    ):
        score -= 4

    return max(
        60,
        score
    )


def measurability(question):

    score = 94

    q = normalize(question)

    all_verbs = []

    for values in BLOOM_VERBS.values():
        all_verbs.extend(values)

    if any(
        re.search(
            r"\b" + re.escape(v) + r"\b",
            q
        )
        for v in all_verbs
    ):
        score += 4

    if "your thoughts" in q:
        score -= 5

    if "what do you think" in q:
        score -= 5

    return max(
        70,
        min(100, score)
    )


def evaluate_question(
    question,
    clos,
    plos,
    target_bloom
):

    best_clo, clo = best_match(
        question,
        clos
    )

    best_plo, plo = best_match(
        question,
        plos
    )

    bloom = bloom_score(
        question,
        target_bloom
    )

    relevance = round(
        (clo + plo) / 2
    )

    clarity_score = clarity(
        question
    )

    measure_score = measurability(
        question
    )

    overall = round(
        clo * 0.20
        + plo * 0.15
        + bloom * 0.20
        + relevance * 0.15
        + clarity_score * 0.15
        + measure_score * 0.15
    )

    return {
        "CLO Match": clo,
        "PLO Match": plo,
        "Bloom": bloom,
        "Relevance": relevance,
        "Clarity": clarity_score,
        "Measurability": measure_score,
        "Overall": overall,
        "Best CLO": best_clo,
        "Best PLO": best_plo,
        "Detected Bloom": detect_bloom(question),
        "Question Type": get_question_type(question)
    }


def status(score):

    if score >= 85:
        return "Strong"

    if score >= 75:
        return "Attained"

    if score >= 65:
        return "Minor Revision"

    if score >= 50:
        return "Review"

    return "Needs Revision"


# ============================================================
# REVISION
# ============================================================

def outcome_content(outcome):

    if not outcome:
        return ""

    text = clean_text(outcome)

    for level in BLOOM_LEVELS:

        for verb in BLOOM_VERBS[level]:

            pattern = (
                r"^"
                + re.escape(verb)
                + r"\b[\s,:-]*"
            )

            if re.search(
                pattern,
                text,
                re.I
            ):

                return clean_text(
                    re.sub(
                        pattern,
                        "",
                        text,
                        flags=re.I
                    )
                )

    return text


def get_options(question):

    lines = question.splitlines()

    options = []

    for line in lines:

        line = clean_text(line)

        if re.match(
            r"^[A-Da-d][\)\.\:]\s+",
            line
        ):
            options.append(line)

    return options


def get_stem(question):

    lines = question.splitlines()

    result = []

    for line in lines:

        line = clean_text(line)

        if re.match(
            r"^[A-Da-d][\)\.\:]\s+",
            line
        ):
            continue

        result.append(line)

    return clean_text(
        "\n".join(result)
    )


def create_revision(
    question,
    result,
    target_bloom
):

    content = outcome_content(
        result["Best CLO"]
    )

    if not content:
        content = outcome_content(
            result["Best PLO"]
        )

    if not content:
        return question

    qtype = result[
        "Question Type"
    ]

    # --------------------------------------------------------
    # MCQ
    # --------------------------------------------------------

    if qtype == "MCQ":

        options = get_options(
            question
        )

        if target_bloom == "Remember":
            stem = (
                f"Which statement correctly identifies "
                f"{content}?"
            )

        elif target_bloom == "Understand":
            stem = (
                f"Which statement best explains "
                f"{content}?"
            )

        elif target_bloom == "Apply":
            stem = (
                f"Which option correctly applies "
                f"the relevant principles of {content}?"
            )

        elif target_bloom == "Analyze":
            stem = (
                f"Which option best analyzes "
                f"{content}?"
            )

        elif target_bloom == "Evaluate":
            stem = (
                f"Which option provides the most appropriate "
                f"evaluation of {content}?"
            )

        else:
            stem = (
                f"Which option represents an appropriate "
                f"solution for {content}?"
            )

        if options:
            return (
                stem
                + "\n"
                + "\n".join(options)
            )

        return stem

    # --------------------------------------------------------
    # Numerical
    # --------------------------------------------------------

    if qtype == "Numerical / Problem Solving":

        stem = get_stem(question)

        return (
            stem.rstrip(".? ")
            + f". Show the calculation steps and state the "
            f"final answer with the appropriate unit."
        )

    # --------------------------------------------------------
    # True / False
    # --------------------------------------------------------

    if qtype == "True / False":

        return (
            f"True or False: {content}."
        )

    # --------------------------------------------------------
    # Fill blank
    # --------------------------------------------------------

    if qtype == "Fill in the Blank":

        original = get_stem(question)

        if "_" in original:

            prefix = original.split("_")[0]

            return (
                prefix.rstrip(": .?")
                + ": ________."
            )

        return (
            f"Complete the statement: "
            f"{content}: ________."
        )

    # --------------------------------------------------------
    # Case
    # --------------------------------------------------------

    if qtype == "Case Study / Application":

        original = get_stem(question)

        if target_bloom == "Analyze":

            task = (
                f"Analyze the case and identify the main "
                f"issues related to {content}."
            )

        elif target_bloom == "Evaluate":

            task = (
                f"Evaluate the case in relation to {content} "
                f"and justify your conclusion."
            )

        elif target_bloom == "Apply":

            task = (
                f"Apply the relevant principles of {content} "
                f"to the case and determine an appropriate response."
            )

        else:

            task = (
                f"Using the case information, explain {content}."
            )

        return (
            original.rstrip(".? ")
            + ". "
            + task
        )

    # --------------------------------------------------------
    # General short / essay
    # --------------------------------------------------------

    if target_bloom == "Remember":

        return (
            f"Identify and state the key points about "
            f"{content}."
        )

    if target_bloom == "Understand":

        return (
            f"Explain {content}."
        )

    if target_bloom == "Apply":

        return (
            f"Apply the relevant principles of "
            f"{content} to solve the given problem or situation."
        )

    if target_bloom == "Analyze":

        return (
            f"Analyze {content}."
        )

    if target_bloom == "Evaluate":

        return (
            f"Evaluate {content} and justify your conclusion."
        )

    if target_bloom == "Create":

        return (
            f"Design or develop an appropriate solution "
            f"for {content}."
        )

    return question


# ============================================================
# REVISION VALIDATION
# ============================================================

def valid_revision(
    original,
    revised
):

    if not revised:
        return False

    if normalize(original) == normalize(revised):
        return False

    forbidden = [
        "clo",
        "plo",
        "learning outcome",
        "course outcome",
        "according to the",
        "as stated in"
    ]

    rev = normalize(revised)

    for phrase in forbidden:

        if phrase in rev:
            return False

    if extract_numbers(
        original
    ) != extract_numbers(
        revised
    ):
        return False

    original_options = [
        normalize(x)
        for x in get_options(original)
    ]

    revised_options = [
        normalize(x)
        for x in get_options(revised)
    ]

    if original_options:

        if original_options != revised_options:
            return False

    return True


# ============================================================
# SESSION STATE
# ============================================================

if "text" not in st.session_state:
    st.session_state.text = ""

if "questions" not in st.session_state:
    st.session_state.questions = []

if "results" not in st.session_state:
    st.session_state.results = []

if "revisions" not in st.session_state:
    st.session_state.revisions = {}

# ============================================================
# HEADER
# ============================================================

st.title("🎓 OBE Quiz Checker")

st.write(
    "Check assessment questions for CLO, PLO, Bloom's Taxonomy, "
    "clarity, relevance, and measurability."
)

# ============================================================
# INFORMATION
# ============================================================

st.subheader("1. Assessment Information")

c1, c2, c3 = st.columns(3)

with c1:
    course = st.text_input(
        "Course / Subject"
    )

with c2:
    assessment = st.text_input(
        "Assessment Name"
    )

with c3:
    target_bloom = st.selectbox(
        "Target Bloom's Level",
        BLOOM_LEVELS,
        index=1
    )

# ============================================================
# OUTCOMES
# ============================================================

st.subheader("2. Learning Outcomes")

clo_input = st.text_area(
    "CLOs",
    height=130,
    placeholder=(
        "CLO 1: Explain the process of photosynthesis and "
        "its importance to plant growth.\n"
        "CLO 2: Analyze factors affecting plant growth."
    )
)

plo_input = st.text_area(
    "PLOs",
    height=130,
    placeholder=(
        "PLO 1: Apply knowledge of the discipline to solve problems.\n"
        "PLO 2: Analyze and communicate solutions effectively."
    )
)

clos = parse_outcomes(
    clo_input
)

plos = parse_outcomes(
    plo_input
)

# ============================================================
# FILE UPLOAD
# ============================================================

st.subheader("3. Upload Complete Assessment")

uploaded = st.file_uploader(
    "Upload PDF, Word, PowerPoint, Excel, CSV, text file, or image",
    type=[
        "pdf",
        "docx",
        "pptx",
        "xlsx",
        "xls",
        "csv",
        "txt",
        "md",
        "png",
        "jpg",
        "jpeg",
        "webp",
        "bmp",
        "tif",
        "tiff"
    ]
)

if uploaded:

    st.write(
        f"**Selected file:** {uploaded.name}"
    )

    if st.button(
        "📖 Read File",
        use_container_width=True
    ):

        with st.spinner(
            "Reading your assessment..."
        ):

            try:

                text = read_file(
                    uploaded
                )

                st.session_state.text = text

                questions = extract_questions(
                    text
                )

                st.session_state.questions = questions

                if text:

                    st.success(
                        f"File successfully read. "
                        f"{len(questions)} question(s) detected."
                    )

                    with st.expander(
                        "View extracted text"
                    ):
                        st.text(
                            text[:20000]
                        )

                    if not questions:

                        st.warning(
                            "The file is readable, but the question structure "
                            "could not be detected automatically."
                        )

                        st.info(
                            "The extracted text is displayed above so you can "
                            "verify what the system received."
                        )

                else:

                    st.error(
                        "The file was opened but contained no readable text."
                    )

            except Exception as error:

                st.error(
                    "The file could not be read."
                )

                st.write(
                    f"**Technical detail:** {error}"
                )

                st.info(
                    "For PDF files, make sure PyMuPDF, pypdf, and pdfplumber "
                    "are installed in requirements.txt."
                )

# ============================================================
# SHOW QUESTIONS
# ============================================================

if st.session_state.questions:

    st.success(
        f"{len(st.session_state.questions)} assessment question(s) ready for analysis."
    )

    with st.expander(
        "View detected questions"
    ):

        for i, q in enumerate(
            st.session_state.questions,
            start=1
        ):

            st.markdown(
                f"**Question {i}**"
            )

            st.write(q)

            st.divider()

# ============================================================
# ANALYZE
# ============================================================

st.subheader("4. Analyze Assessment")

if st.button(
    "🔍 Analyze Assessment",
    type="primary",
    use_container_width=True
):

    if not st.session_state.text:

        st.error(
            "Please upload and read the assessment file first."
        )

        st.stop()

    if not st.session_state.questions:

        st.error(
            "Readable text was found, but no assessment questions "
            "could be identified."
        )

        st.stop()

    results = []

    progress = st.progress(0)

    total = len(
        st.session_state.questions
    )

    for i, question in enumerate(
        st.session_state.questions
    ):

        result = evaluate_question(
            question,
            clos,
            plos,
            target_bloom
        )

        result["Number"] = i + 1
        result["Question"] = question
        result["Status"] = status(
            result["Overall"]
        )

        results.append(result)

        progress.progress(
            (i + 1) / total
        )

    st.session_state.results = results

    revisions = {}

    for result in results:

        weak = (
            result["Overall"] < ATTAINMENT
            or result["CLO Match"] < ATTAINMENT
            or result["PLO Match"] < ATTAINMENT
            or result["Bloom"] < ATTAINMENT
        )

        if not weak:
            continue

        revision = create_revision(
            result["Question"],
            result,
            target_bloom
        )

        if valid_revision(
            result["Question"],
            revision
        ):

            revisions[
                result["Number"]
            ] = revision

    st.session_state.revisions = revisions

    st.success(
        "Assessment analysis completed."
    )

# ============================================================
# RESULTS
# ============================================================

if st.session_state.results:

    results = st.session_state.results

    st.subheader("5. Overall Alignment")

    overall = round(
        sum(
            r["Overall"]
            for r in results
        ) / len(results)
    )

    a, b, c, d = st.columns(4)

    with a:
        st.metric(
            "Overall Score",
            f"{overall}%"
        )

    with b:
        st.metric(
            "Questions",
            len(results)
        )

    with c:
        st.metric(
            "Attained",
            sum(
                r["Overall"] >= 75
                for r in results
            )
        )

    with d:
        st.metric(
            "Needs Revision",
            sum(
                r["Overall"] < 75
                for r in results
            )
        )

    if overall >= 80:

        st.success(
            "🏆 Overall assessment alignment is 80% or above."
        )

        st.balloons()

    # ========================================================
    # REVISION LIST
    # ========================================================

    st.subheader(
        "6. Questions Requiring Revision"
    )

    weak = [
        r for r in results
        if r["Number"]
        in st.session_state.revisions
    ]

    if not weak:

        st.success(
            "🎉 No question requires revision."
        )

    for result in weak:

        number = result["Number"]

        st.markdown(
            f"### Question {number} — "
            f"{result['Overall']}%"
        )

        st.write(
            "**Current Question:**"
        )

        st.info(
            result["Question"]
        )

        st.write(
            "**Problem Identified:**"
        )

        if result["CLO Match"] < 75:
            st.write(
                "• CLO alignment is below 75%."
            )

        if result["PLO Match"] < 75:
            st.write(
                "• PLO alignment is below 75%."
            )

        if result["Bloom"] < 75:
            st.write(
                "• Bloom's level does not sufficiently match the target level."
            )

        if result["Clarity"] < 75:
            st.write(
                "• Wording needs greater clarity."
            )

        if result["Measurability"] < 75:
            st.write(
                "• The task needs to be more measurable."
            )

        st.write(
            "**Practical Revision:**"
        )

        revision = st.session_state.revisions[
            number
        ]

        st.success(
            revision
        )

        if st.button(
            f"✅ Use This Revision — Question {number}",
            key=f"revision_{number}",
            use_container_width=True
        ):

            # Replace question
            st.session_state.questions[
                number - 1
            ] = revision

            # Rescore
            new_result = evaluate_question(
                revision,
                clos,
                plos,
                target_bloom
            )

            new_result["Number"] = number
            new_result["Question"] = revision
            new_result["Status"] = status(
                new_result["Overall"]
            )

            # Replace old result
            for i, old in enumerate(
                st.session_state.results
            ):

                if old["Number"] == number:

                    st.session_state.results[
                        i
                    ] = new_result

                    break

            # Remove revision if attained
            if new_result["Overall"] >= 75:

                st.session_state.revisions.pop(
                    number,
                    None
                )

            st.rerun()

        st.divider()

    # ========================================================
    # ATTAINED
    # ========================================================

    st.subheader(
        "7. Attained Questions"
    )

    attained = [
        r for r in st.session_state.results
        if r["Overall"] >= 75
    ]

    if attained:

        attained_df = pd.DataFrame([
            {
                "Question": r["Number"],
                "Score": f"{r['Overall']}%",
                "Status": r["Status"],
                "Type": r["Question Type"]
            }
            for r in attained
        ])

        st.dataframe(
            attained_df,
            use_container_width=True,
            hide_index=True
        )

    else:

        st.info(
            "No questions have reached 75% yet."
        )

    # ========================================================
    # ONE GRAPH ONLY
    # ========================================================

    st.subheader(
        "8. Alignment Overview"
    )

    graph = pd.DataFrame([
        {
            "Question": f"Q{r['Number']}",
            "Score": r["Overall"]
        }
        for r in st.session_state.results
    ])

    st.bar_chart(
        graph.set_index(
            "Question"
        )
    )

    # ========================================================
    # QUESTION OVERVIEW
    # ========================================================

    st.subheader(
        "9. Question Overview"
    )

    overview = pd.DataFrame([
        {
            "Question": r["Number"],
            "Type": r["Question Type"],
            "Score": f"{r['Overall']}%",
            "Status": r["Status"],
            "CLO": f"{r['CLO Match']}%",
            "PLO": f"{r['PLO Match']}%",
            "Bloom": f"{r['Bloom']}%",
            "Clarity": f"{r['Clarity']}%",
            "Measurability": f"{r['Measurability']}%"
        }
        for r in results
    ])

    st.dataframe(
        overview,
        use_container_width=True,
        hide_index=True
    )

    # ========================================================
    # DETAILS
    # ========================================================

    st.subheader(
        "10. Detailed Question Analysis"
    )

    selected = st.selectbox(
        "Select Question",
        [
            r["Number"]
            for r in results
        ]
    )

    result = next(
        r for r in results
        if r["Number"] == selected
    )

    st.write(
        f"**Question {selected}**"
    )

    st.info(
        result["Question"]
    )

    details = pd.DataFrame([
        {
            "Metric": "Overall",
            "Score": f"{result['Overall']}%"
        },
        {
            "Metric": "CLO Match",
            "Score": f"{result['CLO Match']}%"
        },
        {
            "Metric": "PLO Match",
            "Score": f"{result['PLO Match']}%"
        },
        {
            "Metric": "Bloom",
            "Score": f"{result['Bloom']}%"
        },
        {
            "Metric": "Relevance",
            "Score": f"{result['Relevance']}%"
        },
        {
            "Metric": "Clarity",
            "Score": f"{result['Clarity']}%"
        },
        {
            "Metric": "Measurability",
            "Score": f"{result['Measurability']}%"
        }
    ])

    st.dataframe(
        details,
        use_container_width=True,
        hide_index=True
    )

    if result["Best CLO"]:

        st.write(
            "**Best Matching CLO:**"
        )

        st.write(
            result["Best CLO"]
        )

    if result["Best PLO"]:

        st.write(
            "**Best Matching PLO:**"
        )

        st.write(
            result["Best PLO"]
        )

    # ========================================================
    # EXPORT
    # ========================================================

    st.subheader(
        "11. Export"
    )

    export_rows = []

    for r in st.session_state.results:

        export_rows.append({
            "Question No": r["Number"],
            "Question": r["Question"],
            "Question Type": r["Question Type"],
            "Overall Score": r["Overall"],
            "Status": r["Status"],
            "CLO Match": r["CLO Match"],
            "PLO Match": r["PLO Match"],
            "Bloom": r["Bloom"],
            "Relevance": r["Relevance"],
            "Clarity": r["Clarity"],
            "Measurability": r["Measurability"],
            "Detected Bloom": r["Detected Bloom"],
            "Best CLO": r["Best CLO"],
            "Best PLO": r["Best PLO"],
            "Suggested Revision":
                st.session_state.revisions.get(
                    r["Number"],
                    ""
                )
        })

    export_df = pd.DataFrame(
        export_rows
    )

    csv = export_df.to_csv(
        index=False
    ).encode("utf-8")

    st.download_button(
        "⬇️ Download Report",
        data=csv,
        file_name="obe_quiz_checker_report.csv",
        mime="text/csv",
        use_container_width=True
    )

# ============================================================
# FOOTER
# ============================================================

st.divider()

st.caption(
    "OBE Quiz Checker | PDF • DOCX • PPTX • XLSX • CSV • TXT • Images"
)
